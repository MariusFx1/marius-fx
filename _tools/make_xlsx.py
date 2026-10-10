"""Generează downloads/jurnal-tranzactionare-marius-fx.xlsx (șablon gol pentru elevi)."""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

OUT = "/workspace/forex-ms/downloads/jurnal-tranzactionare-marius-fx.xlsx"
NAVY, NAVY2, GREEN, BAND, GRID, MUTED = "0B1F3A", "14294A", "16A34A", "F4F7FB", "D5DCE6", "5B6B82"
FONT = "Calibri"
thin = Side(style="thin", color=GRID)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
TITLE_FILL = PatternFill("solid", fgColor=NAVY)
BAND_FILL = PatternFill("solid", fgColor=BAND)
CALC_FILL = PatternFill("solid", fgColor="EAF6EE")   # coloane calculate automat
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10.5)

FIRST, N_ROWS = 4, 100
LAST = FIRST + N_ROWS - 1

# (antet, lățime, tip)
COLS = [
    ("Nr.", 6, "nr"),
    ("Data", 12, "date"),
    ("Pereche", 11, "text"),
    ("Direcție", 10, "list"),
    ("Sesiune", 12, "list"),
    ("Setup / motiv intrare", 32, "wrap"),
    ("Preț intrare", 12, "price"),
    ("Stop loss", 12, "price"),
    ("Take profit", 12, "price"),
    ("Lot", 8, "lot"),
    ("Risc %", 8, "pct"),
    ("Risc în bani", 12, "money"),
    ("R:R planificat", 11, "rr"),
    ("Rezultat", 11, "list"),
    ("Rezultat în R", 11, "r"),
    ("Am respectat planul?", 12, "list"),
    ("Emoții", 20, "wrap"),
    ("Lecție / ce îmbunătățesc", 36, "wrap"),
]
NCOL = len(COLS)
L = {h: get_column_letter(i + 1) for i, (h, _, _) in enumerate(COLS)}

wb = Workbook()
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.properties.creator = "MS Prime"
wb.properties.title = "Jurnal de tranzacționare — MS Prime"
wb.properties.subject = "Șablon educațional de jurnal de tranzacționare"

# ---------------------------------------------------------------- Jurnal
ws = wb.active
ws.title = "Jurnal"
ws.sheet_properties.tabColor = NAVY
ws.sheet_view.showGridLines = False
ws.sheet_view.zoomScale = 100

last_col = get_column_letter(NCOL)
ws.merge_cells(f"A1:{last_col}1")
ws["A1"] = "Jurnal de tranzacționare — MS Prime"
ws["A1"].font = Font(name=FONT, bold=True, size=16, color="FFFFFF")
ws["A1"].alignment = Alignment(vertical="center", indent=1)
ws.merge_cells(f"A2:{last_col}2")
ws["A2"] = ("Completează un rând pentru fiecare tranzacție. Coloanele verzi se calculează automat. "
            "Vezi foaia „Cum completezi” pentru explicații și un exemplu. "
            "Conținut educațional, nu consultanță financiară.")
ws["A2"].font = Font(name=FONT, italic=True, size=9.5, color=MUTED)
ws["A2"].alignment = Alignment(vertical="center", indent=1)
for c in range(1, NCOL + 1):
    ws.cell(row=1, column=c).fill = TITLE_FILL
ws.row_dimensions[1].height = 34
ws.row_dimensions[2].height = 22
ws.row_dimensions[3].height = 36

for i, (head, width, kind) in enumerate(COLS, start=1):
    col = get_column_letter(i)
    ws.column_dimensions[col].width = width
    cell = ws.cell(row=3, column=i, value=head)
    cell.font = HEAD_FONT
    cell.fill = HEAD_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(left=Side(style="thin", color=NAVY2), right=Side(style="thin", color=NAVY2),
                         top=Side(style="thin", color=NAVY), bottom=Side(style="medium", color=GREEN))

FMT = {
    "date": "DD.MM.YYYY", "price": "0.00000", "lot": "0.00", "pct": '0.0"%"',
    "money": "#,##0.00", "rr": '"1:"0.0#', "r": '+0.00"R";-0.00"R";0.00"R"', "nr": "0",
}
CENTER = {"nr", "date", "text", "list", "lot", "pct", "rr", "r"}

for r in range(FIRST, LAST + 1):
    ws.row_dimensions[r].height = 21
    band = (r - FIRST) % 2 == 1
    for i, (head, _, kind) in enumerate(COLS, start=1):
        cell = ws.cell(row=r, column=i)
        cell.border = BORDER
        cell.font = Font(name=FONT, size=10.5, color="1F2A3C")
        if kind in FMT:
            cell.number_format = FMT[kind]
        cell.alignment = Alignment(horizontal="center" if kind in CENTER else ("right" if kind in {"price", "money"} else "left"),
                                   vertical="center", wrap_text=(kind == "wrap"))
        if kind in {"rr", "r", "nr"}:
            cell.fill = CALC_FILL
        elif band:
            cell.fill = BAND_FILL
    D, G, H, I_, M, N = L["Data"], L["Preț intrare"], L["Stop loss"], L["Take profit"], L["R:R planificat"], L["Rezultat"]
    ws[f"A{r}"] = f'=IF({D}{r}="","",ROW()-{FIRST - 1})'
    ws[f"{M}{r}"] = f'=IF(OR({G}{r}="",{H}{r}="",{I_}{r}=""),"",IFERROR(ROUND(ABS({I_}{r}-{G}{r})/ABS({G}{r}-{H}{r}),2),""))'
    ws[f"{L['Rezultat în R']}{r}"] = (f'=IF({N}{r}="TP",IF({M}{r}="","",{M}{r}),'
                                      f'IF({N}{r}="SL",-1,IF({N}{r}="BE",0,"")))')
    ws[f"{L['Rezultat în R']}{r}"].font = Font(name=FONT, size=10.5, bold=True, color="1F2A3C")

ws.freeze_panes = f"C{FIRST}"       # antetul + Nr./Data rămân vizibile
ws.auto_filter.ref = f"A3:{last_col}{LAST}"


def add_list(sheet, col, items, title, prompt, rng=None):
    dv = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True,
                        showDropDown=False, showErrorMessage=True, showInputMessage=True,
                        errorTitle="Valoare nevalidă", error="Alege o valoare din listă: " + ", ".join(items) + ".",
                        promptTitle=title, prompt=prompt)
    sheet.add_data_validation(dv)
    dv.add(rng or f"{col}{FIRST}:{col}{LAST}")
    return dv


add_list(ws, L["Direcție"], ["Buy", "Sell"], "Direcție", "Buy = cumperi, Sell = vinzi.")
add_list(ws, L["Sesiune"], ["Sydney", "Tokyo", "Londra", "New York"], "Sesiune", "Sesiunea în care ai intrat.")
add_list(ws, L["Rezultat"], ["TP", "SL", "BE", "Manual"], "Rezultat",
         "TP = take profit, SL = stop loss, BE = break-even, Manual = închis de tine (scrie singur Rezultatul în R).")
add_list(ws, L["Am respectat planul?"], ["Da", "Nu"], "Planul", "Ai respectat 100% regulile planului?")


def add_num(sheet, col, kind, lo, hi, title, prompt, err):
    dv = DataValidation(type=kind, operator="between", formula1=str(lo), formula2=str(hi), allow_blank=True,
                        showErrorMessage=True, showInputMessage=True, errorTitle="Valoare nevalidă", error=err,
                        promptTitle=title, prompt=prompt)
    sheet.add_data_validation(dv)
    dv.add(f"{col}{FIRST}:{col}{LAST}")


d0, d1 = (dt.date(2000, 1, 1) - dt.date(1899, 12, 30)).days, (dt.date(2100, 12, 31) - dt.date(1899, 12, 30)).days
add_num(ws, L["Data"], "date", d0, d1, "Data", "Ex.: 12.01.2026", "Introdu o dată validă (ex.: 12.01.2026).")
for h in ("Preț intrare", "Stop loss", "Take profit"):
    add_num(ws, L[h], "decimal", 0, 1000000, h, "Prețul exact din platformă.", "Introdu un număr pozitiv.")
add_num(ws, L["Lot"], "decimal", 0, 1000, "Lot", "Mărimea poziției (ex.: 0.10).", "Introdu un număr pozitiv (ex.: 0.10).")
add_num(ws, L["Risc %"], "decimal", 0, 100, "Risc %", "Scrie doar numărul: 1 = 1% din cont. Ideal 1–2%.",
        "Introdu un procent între 0 și 100 (ex.: 1).")
add_num(ws, L["Risc în bani"], "decimal", 0, 100000000, "Risc în bani", "Suma pierdută dacă se atinge stop loss-ul.",
        "Introdu o sumă pozitivă.")

# formatare condiționată
rng_r = f"{L['Rezultat în R']}{FIRST}:{L['Rezultat în R']}{LAST}"
ws.conditional_formatting.add(rng_r, FormulaRule(formula=[f'AND(ISNUMBER({L["Rezultat în R"]}{FIRST}),{L["Rezultat în R"]}{FIRST}>0)'],
                                                 font=Font(color="15803D", bold=True)))
ws.conditional_formatting.add(rng_r, FormulaRule(formula=[f'AND(ISNUMBER({L["Rezultat în R"]}{FIRST}),{L["Rezultat în R"]}{FIRST}<0)'],
                                                 font=Font(color="B91C1C", bold=True)))
P = L["Am respectat planul?"]
ws.conditional_formatting.add(f"{P}{FIRST}:{P}{LAST}", CellIsRule(operator="equal", formula=['"Nu"'],
                              fill=PatternFill("solid", fgColor="FDE2E2"), font=Font(color="B91C1C", bold=True)))
ws.conditional_formatting.add(f"{P}{FIRST}:{P}{LAST}", CellIsRule(operator="equal", formula=['"Da"'],
                              fill=PatternFill("solid", fgColor="DCFCE7"), font=Font(color="15803D", bold=True)))

ws.page_setup.orientation = "landscape"
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = "3:3"
ws.page_margins.left = ws.page_margins.right = 0.4
ws.oddFooter.center.text = "MS Prime — Jurnal de tranzacționare · pagina &P din &N"

# ---------------------------------------------------------------- Rezumat
rz = wb.create_sheet("Rezumat")
rz.sheet_properties.tabColor = GREEN
rz.sheet_view.showGridLines = False
rz.column_dimensions["A"].width = 3
rz.column_dimensions["B"].width = 38
rz.column_dimensions["C"].width = 16
rz.column_dimensions["D"].width = 58
rz.merge_cells("B2:D2")
rz["B2"] = "Rezumat — se calculează automat din foaia „Jurnal”"
rz["B2"].font = Font(name=FONT, bold=True, size=15, color="FFFFFF")
rz["B2"].alignment = Alignment(vertical="center", indent=1)
for c in "BCD":
    rz[f"{c}2"].fill = TITLE_FILL
rz.row_dimensions[2].height = 32
for c, t in zip("BCD", ("Indicator", "Valoare", "Cum se calculează")):
    cell = rz[f"{c}4"]
    cell.value = t
    cell.font = HEAD_FONT
    cell.fill = HEAD_FILL
    cell.alignment = Alignment(horizontal="center" if c == "C" else "left", vertical="center", indent=0 if c == "C" else 1)
    cell.border = Border(bottom=Side(style="medium", color=GREEN))
rz.row_dimensions[4].height = 26

J = "Jurnal!"
rN = f"{J}${L['Rezultat']}${FIRST}:${L['Rezultat']}${LAST}"
rR = f"{J}${L['Rezultat în R']}${FIRST}:${L['Rezultat în R']}${LAST}"
rP = f"{J}${P}${FIRST}:${P}${LAST}"
rows = [
    ("Total tranzacții", f"=COUNTIF({rN},\"?*\")", "0", "Rânduri cu „Rezultat” completat."),
    ("Câștigate", f"=COUNTIF({rR},\">0\")", "0", "Tranzacții cu Rezultat în R peste 0."),
    ("Pierdute", f"=COUNTIF({rR},\"<0\")", "0", "Tranzacții cu Rezultat în R sub 0."),
    ("Break-even", f"=COUNTIF({rN},\"BE\")", "0", "Tranzacții închise la break-even (0R)."),
    ("Rata de câștig %", "=IF(C5=0,\"—\",C6/C5)", "0.0%", "Câștigate ÷ total tranzacții."),
    ("Total R", f"=SUM({rR})", '+0.00"R";-0.00"R";0.00"R"', "Suma coloanei „Rezultat în R”."),
    ("R mediu pe tranzacție", f"=IF(COUNT({rR})=0,\"—\",AVERAGE({rR}))", '+0.00"R";-0.00"R";0.00"R"',
     "Media coloanei „Rezultat în R”."),
    ("% tranzacții cu planul respectat", f"=IF(COUNTIF({rP},\"Da\")+COUNTIF({rP},\"Nu\")=0,\"—\","
     f"COUNTIF({rP},\"Da\")/(COUNTIF({rP},\"Da\")+COUNTIF({rP},\"Nu\")))", "0.0%",
     "Răspunsuri „Da” ÷ (Da + Nu). Ținta ta: cât mai aproape de 100%."),
]
for k, (label, formula, fmt, how) in enumerate(rows):
    r = 5 + k
    rz.row_dimensions[r].height = 24
    fill = BAND_FILL if k % 2 else PatternFill(fill_type=None)
    for c, v in zip("BCD", (label, formula, how)):
        cell = rz[f"{c}{r}"]
        cell.value = v
        cell.border = Border(bottom=thin)
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", horizontal="center" if c == "C" else "left", indent=0 if c == "C" else 1)
        cell.font = Font(name=FONT, size=11, bold=(c != "D"), color=MUTED if c == "D" else "1F2A3C")
    rz[f"C{r}"].number_format = fmt
rz["C5"].font = Font(name=FONT, size=12, bold=True, color=NAVY)
rz.conditional_formatting.add("C10:C11", CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="15803D", bold=True)))
rz.conditional_formatting.add("C10:C11", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="B91C1C", bold=True)))
rz.merge_cells("B14:D15")
rz["B14"] = ("Notă: un rezultat pozitiv pe câteva tranzacții nu garantează nimic. Urmărește mai ales disciplina "
             "(% plan respectat) și lecțiile notate. Conținut educațional, nu consultanță financiară.")
rz["B14"].font = Font(name=FONT, italic=True, size=9.5, color=MUTED)
rz["B14"].alignment = Alignment(wrap_text=True, vertical="top", indent=1)

# ---------------------------------------------------------------- Cum completezi
hw = wb.create_sheet("Cum completezi")
hw.sheet_properties.tabColor = "64748B"
hw.sheet_view.showGridLines = False
hw.column_dimensions["A"].width = 3
hw.column_dimensions["B"].width = 6
hw.merge_cells("B2:J2")
hw["B2"] = "Cum completezi jurnalul"
hw["B2"].font = Font(name=FONT, bold=True, size=15, color="FFFFFF")
hw["B2"].alignment = Alignment(vertical="center", indent=1)
for c in "BCDEFGHIJ":
    hw[f"{c}2"].fill = TITLE_FILL
hw.row_dimensions[2].height = 32
steps = [
    "Completează câte un rând în foaia „Jurnal” imediat după ce intri în tranzacție (plan) și după ce o închizi (rezultat).",
    "Data, Pereche, Direcție și Sesiune: alege Direcția (Buy/Sell) și Sesiunea din listele derulante.",
    "Setup / motiv intrare: scrie de ce ai intrat (ex.: retest de suport pe H4 + confirmare). Dacă nu poți scrie motivul, nu intra.",
    "Preț intrare, Stop loss, Take profit: copiază prețurile exacte din platformă. „R:R planificat” se calculează singur.",
    "Lot, Risc %, Risc în bani: notează mărimea poziției și cât pierzi dacă se atinge stop loss-ul. La Risc % scrie doar numărul (1 = 1%).",
    "Rezultat: TP, SL, BE sau Manual. „Rezultat în R” se completează automat pentru TP (= R:R planificat), SL (−1R) și BE (0R).",
    "Dacă ai închis manual, scrie tu „Rezultat în R” peste formulă (ex.: 0,6 sau −0,4).",
    "Am respectat planul? Răspunde sincer cu Da sau Nu — este cel mai important indicator din jurnal.",
    "Emoții și Lecție: ce ai simțit (calm, grabă, frică, FOMO) și ce faci diferit data viitoare.",
    "La finalul săptămânii, deschide foaia „Rezumat” și recitește lecțiile. Caută greșelile care se repetă.",
]
for k, s in enumerate(steps):
    r = 4 + k
    hw[f"B{r}"] = k + 1
    hw[f"B{r}"].font = Font(name=FONT, bold=True, size=11, color=GREEN)
    hw[f"B{r}"].alignment = Alignment(horizontal="center", vertical="top")
    hw.merge_cells(f"C{r}:J{r}")
    hw[f"C{r}"] = s
    hw[f"C{r}"].font = Font(name=FONT, size=11, color="1F2A3C")
    hw[f"C{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    hw.row_dimensions[r].height = 30

# rând exemplu (pe aceeași structură de coloane, începând din coloana B)
ex_title = 4 + len(steps) + 1          # 15
ex_head, ex_row = ex_title + 1, ex_title + 2
for i, (head, width, kind) in enumerate(COLS):
    col = get_column_letter(i + 2)
    if i + 2 > 2:
        hw.column_dimensions[col].width = width
hw.merge_cells(start_row=ex_title, start_column=2, end_row=ex_title, end_column=NCOL + 1)
t = hw.cell(row=ex_title, column=2, value="Exemplu — date fictive pentru demonstrație (nu este o tranzacție reală)")
t.font = Font(name=FONT, bold=True, size=11.5, color="92400E")
t.fill = PatternFill("solid", fgColor="FEF3C7")
t.alignment = Alignment(vertical="center", indent=1)
hw.row_dimensions[ex_title].height = 26
example = [1, dt.date(2026, 1, 12), "EURUSD", "Buy", "Londra", "Retest suport H4 + pin bar de confirmare",
           1.08500, 1.08300, 1.08900, 0.10, 1, 20.00, None, "TP", None, "Da", "Calm, am avut răbdare",
           "Am așteptat confirmarea, nu am intrat devreme. Păstrez regula."]
for i, ((head, _, kind), v) in enumerate(zip(COLS, example)):
    c = i + 2
    h = hw.cell(row=ex_head, column=c, value=head)
    h.font = HEAD_FONT
    h.fill = HEAD_FILL
    h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    h.border = BORDER
    cell = hw.cell(row=ex_row, column=c, value=v)
    cell.border = BORDER
    cell.font = Font(name=FONT, size=10.5, color="1F2A3C", italic=True)
    cell.fill = PatternFill("solid", fgColor="FFFBEB")
    cell.alignment = Alignment(horizontal="center" if kind in CENTER else "left", vertical="center", wrap_text=True)
    if kind in FMT:
        cell.number_format = FMT[kind]
gc = lambda name: get_column_letter(COLS.index(next(x for x in COLS if x[0] == name)) + 2)
hw[f"{gc('R:R planificat')}{ex_row}"] = (f"=ROUND(ABS({gc('Take profit')}{ex_row}-{gc('Preț intrare')}{ex_row})/"
                                         f"ABS({gc('Preț intrare')}{ex_row}-{gc('Stop loss')}{ex_row}),2)")
hw[f"{gc('Rezultat în R')}{ex_row}"] = f"={gc('R:R planificat')}{ex_row}"
hw.row_dimensions[ex_head].height = 36
hw.row_dimensions[ex_row].height = 46
note = ex_row + 2
hw.merge_cells(start_row=note, start_column=2, end_row=note, end_column=NCOL + 1)
hw.cell(row=note, column=2, value=("În exemplu: stop loss la 20 de pipși și take profit la 40 de pipși → R:R 1:2. "
                                   "Riscul este 1% din cont; tranzacția a atins TP → +2R. Valorile sunt inventate, doar pentru a arăta cum se completează."))
hw.cell(row=note, column=2).font = Font(name=FONT, italic=True, size=10, color=MUTED)
hw.cell(row=note, column=2).alignment = Alignment(wrap_text=True, vertical="top", indent=1)
hw.row_dimensions[note].height = 30

for sh in (rz, hw):
    sh.page_setup.orientation = "landscape"
    sh.page_setup.paperSize = sh.PAPERSIZE_A4
    sh.page_setup.fitToWidth, sh.page_setup.fitToHeight = 1, 0
    sh.sheet_properties.pageSetUpPr.fitToPage = True
    sh.page_margins.left = sh.page_margins.right = 0.4
wb.active = 0
wb.save(OUT)
print("salvat", OUT)
